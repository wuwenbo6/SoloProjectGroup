import os
import json
import numpy as np
import pandas as pd
from datetime import datetime
from typing import Dict, Any, Optional
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    Image, PageBreak, HRFlowable
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList


class ReportGenerator:
    def __init__(self, output_dir: str = "reports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def generate_full_report(self, analysis_results: Dict[str, Any],
                             component_id: str = "COMP-001",
                             inspector: str = "System") -> Dict[str, str]:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        base_filename = f"{component_id}_analysis_{timestamp}"

        reports = {
            'pdf': self._generate_pdf_report(analysis_results, component_id, inspector, base_filename),
            'excel': self._generate_excel_report(analysis_results, component_id, inspector, base_filename),
            'json': self._generate_json_report(analysis_results, component_id, inspector, base_filename)
        }

        return reports

    def _generate_pdf_report(self, analysis_results: Dict[str, Any],
                             component_id: str, inspector: str,
                             base_filename: str) -> str:
        pdf_path = os.path.join(self.output_dir, f"{base_filename}.pdf")
        
        def safe_format(value, fmt=".4f"):
            if value is None:
                return "0.0000"
            try:
                return format(float(value), fmt)
            except (ValueError, TypeError):
                return "0.0000"

        doc = SimpleDocTemplate(
            pdf_path,
            pagesize=A4,
            rightMargin=20*mm,
            leftMargin=20*mm,
            topMargin=20*mm,
            bottomMargin=20*mm
        )

        story = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Title'],
            fontSize=20,
            spaceAfter=10,
            alignment=TA_CENTER,
            textColor=colors.HexColor('#1a5276')
        )

        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            spaceAfter=8,
            spaceBefore=15,
            textColor=colors.HexColor('#2874a6')
        )

        normal_style = ParagraphStyle(
            'CustomNormal',
            parent=styles['Normal'],
            fontSize=10,
            spaceAfter=5,
            leading=14
        )

        story.append(Paragraph("点云分析与安全评估报告", title_style))
        story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor('#1a5276')))
        story.append(Spacer(1, 5*mm))

        info_data = [
            ["组件编号", str(component_id)],
            ["检测人员", str(inspector)],
            ["检测日期", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["报告版本", "v1.0"]
        ]
        info_table = Table(info_data, colWidths=[50*mm, 100*mm])
        info_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#d4e6f1')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(info_table)
        story.append(Spacer(1, 8*mm))

        safety = analysis_results.get('safety_summary', {}) if analysis_results else {}
        safety_level = str(safety.get('safety_level', 'UNKNOWN'))
        safety_score = safety.get('safety_score', 0)
        try:
            safety_score = float(safety_score)
        except (ValueError, TypeError):
            safety_score = 0.0

        status_color = self._get_safety_color(safety_level)
        
        safety_header = f"安全等级: {safety_level} (得分: {safety_score:.1f}/100)"
        safety_para = Paragraph(safety_header, ParagraphStyle(
            'SafetyHeader',
            parent=heading_style,
            textColor=status_color,
            fontSize=16
        ))
        story.append(safety_para)
        story.append(Spacer(1, 3*mm))

        if safety.get('needs_immediate_action', False):
            warning_para = Paragraph(
                "<font color='red'><b>⚠️ 需要立即采取行动!</b></font>",
                ParagraphStyle('Warning', parent=normal_style, alignment=TA_CENTER, fontSize=12)
            )
            story.append(warning_para)
            story.append(Spacer(1, 3*mm))

        story.append(Paragraph("1. 厚度分析结果", heading_style))
        thickness_stats = analysis_results.get('thickness_stats', {}) if analysis_results else {}
        if thickness_stats:
            thickness_data = [
                ["指标", "数值", "单位"],
                ["平均厚度", safe_format(thickness_stats.get('mean_thickness', 0)), "m"],
                ["最小厚度", safe_format(thickness_stats.get('min_thickness', 0)), "m"],
                ["最大厚度", safe_format(thickness_stats.get('max_thickness', 0)), "m"],
                ["中位厚度", safe_format(thickness_stats.get('median_thickness', 0)), "m"],
                ["标准差", safe_format(thickness_stats.get('std_thickness', 0)), "m"],
                ["有效点数", str(thickness_stats.get('valid_points', 0)), ""]
            ]
            thickness_table = Table(thickness_data, colWidths=[50*mm, 50*mm, 30*mm])
            thickness_table.setStyle(self._get_table_style())
            story.append(thickness_table)
        story.append(Spacer(1, 5*mm))

        story.append(Paragraph("2. 磨损分析结果", heading_style))
        wear_metrics = analysis_results.get('wear_metrics', {}) if analysis_results else {}
        if wear_metrics:
            wear_data = [
                ["指标", "数值", "单位"],
                ["磨损区域数量", str(wear_metrics.get('num_wear_regions', 0)), ""],
                ["最大磨损深度", safe_format(abs(wear_metrics.get('max_wear_depth', 0))), "m"],
                ["平均磨损深度", safe_format(abs(wear_metrics.get('mean_wear_depth', 0))), "m"],
                ["磨损面积比例", safe_format(wear_metrics.get('wear_ratio', 0)*100, ".2f"), "%"],
                ["总磨损体积", safe_format(wear_metrics.get('total_wear_volume', 0), ".6f"), "m³"]
            ]
            wear_table = Table(wear_data, colWidths=[50*mm, 50*mm, 30*mm])
            wear_table.setStyle(self._get_table_style())
            story.append(wear_table)
        story.append(Spacer(1, 5*mm))

        regions = wear_metrics.get('region_metrics', [])
        if regions:
            story.append(Paragraph("磨损区域详情:", normal_style))
            region_data = [["区域ID", "点数", "最大深度(m)", "平均深度(m)", "估计体积(m³)"]]
            for region in regions:
                region_data.append([
                    str(region.get('region_id', 'N/A')),
                    str(region.get('num_points', 0)),
                    safe_format(abs(region.get('max_depth', 0))),
                    safe_format(abs(region.get('mean_depth', 0))),
                    safe_format(region.get('volume', 0), ".6f")
                ])
            region_table = Table(region_data, colWidths=[25*mm, 25*mm, 30*mm, 30*mm, 30*mm])
            region_table.setStyle(self._get_table_style(True))
            story.append(region_table)
        story.append(Spacer(1, 5*mm))

        story.append(Paragraph("3. 评估建议", heading_style))
        recommendations = safety.get('recommendations', [])
        if recommendations and len(recommendations) > 0:
            for i, rec in enumerate(recommendations, 1):
                story.append(Paragraph(f"{i}. {str(rec)}", normal_style))
        else:
            story.append(Paragraph("无特殊建议", normal_style))
        story.append(Spacer(1, 5*mm))

        key_findings = safety.get('key_findings', [])
        if key_findings and len(key_findings) > 0:
            story.append(Paragraph("<b>主要发现:</b>", normal_style))
            for finding in key_findings:
                story.append(Paragraph(f"• {str(finding)}", normal_style))
        story.append(Spacer(1, 10*mm))

        point_cloud_info = analysis_results.get('point_cloud_info', {}) if analysis_results else {}
        if point_cloud_info:
            story.append(Paragraph("4. 点云数据信息", heading_style))
            info_data = [
                ["指标", "数值"],
                ["总点数", str(point_cloud_info.get('num_points', 0))],
                ["包含法线", "是" if point_cloud_info.get('has_normals', False) else "否"],
                ["包含颜色", "是" if point_cloud_info.get('has_colors', False) else "否"]
            ]
            info_table = Table(info_data, colWidths=[60*mm, 60*mm])
            info_table.setStyle(self._get_table_style())
            story.append(info_table)

        story.append(Spacer(1, 15*mm))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.grey))
        story.append(Spacer(1, 3*mm))
        story.append(Paragraph(
            "本报告由点云分析系统自动生成",
            ParagraphStyle('Footer', parent=normal_style, alignment=TA_CENTER, fontSize=8, textColor=colors.grey)
        ))

        try:
            doc.build(story)
        except Exception as e:
            print(f"PDF生成警告: {e}")
            pass

        return pdf_path

    def _get_table_style(self, has_header: bool = True) -> TableStyle:
        style = [
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]
        if has_header:
            style.append(('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2874a6')))
            style.append(('TEXTCOLOR', (0, 0), (-1, 0), colors.white))
            style.append(('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'))
        return TableStyle(style)

    def _get_safety_color(self, safety_level: str) -> colors.Color:
        color_map = {
            'SAFE': colors.HexColor('#27ae60'),
            'CAUTION': colors.HexColor('#f39c12'),
            'WARNING': colors.HexColor('#e67e22'),
            'CRITICAL': colors.HexColor('#e74c3c'),
            'UNSAFE': colors.HexColor('#c0392b')
        }
        return color_map.get(safety_level, colors.grey)

    def _generate_excel_report(self, analysis_results: Dict[str, Any],
                               component_id: str, inspector: str,
                               base_filename: str) -> str:
        excel_path = os.path.join(self.output_dir, f"{base_filename}.xlsx")
        
        def safe_float(value, default=0.0):
            if value is None:
                return default
            try:
                return float(value)
            except (ValueError, TypeError):
                return default

        def safe_int(value, default=0):
            if value is None:
                return default
            try:
                return int(value)
            except (ValueError, TypeError):
                return default

        try:
            wb = Workbook()
            wb.remove(wb.active)
        except:
            wb = Workbook()
            if 'Sheet' in wb.sheetnames:
                del wb['Sheet']

        header_fill = PatternFill(start_color='2874a6', end_color='2874a6', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        normal_font = Font(size=10)
        center_align = Alignment(horizontal='center', vertical='center')

        ws_summary = wb.create_sheet("报告摘要")
        ws_summary['A1'] = "点云分析与安全评估报告"
        ws_summary['A1'].font = Font(bold=True, size=14, color='1a5276')
        try:
            ws_summary.merge_cells('A1:B1')
        except:
            pass

        summary_data = [
            ["组件编号", str(component_id)],
            ["检测人员", str(inspector)],
            ["检测日期", datetime.now().strftime("%Y-%m-%d %H:%M:%S")]
        ]

        for i, (key, value) in enumerate(summary_data, start=3):
            ws_summary[f'A{i}'] = key
            ws_summary[f'B{i}'] = value
            ws_summary[f'A{i}'].fill = PatternFill(start_color='d4e6f1', end_color='d4e6f1', fill_type='solid')
            ws_summary[f'A{i}'].font = Font(bold=True)

        safety = analysis_results.get('safety_summary', {}) if analysis_results else {}
        ws_summary['A7'] = "安全评估结果"
        ws_summary['A7'].font = Font(bold=True, size=12)
        ws_summary['A8'] = "安全等级"
        ws_summary['B8'] = str(safety.get('safety_level', 'UNKNOWN'))
        ws_summary['A9'] = "安全得分"
        ws_summary['B9'] = f"{safe_float(safety.get('safety_score', 0)):.1f}"
        ws_summary['A10'] = "需要检查"
        ws_summary['B10'] = "是" if safety.get('needs_inspection', False) else "否"
        ws_summary['A11'] = "需要立即行动"
        ws_summary['B11'] = "是" if safety.get('needs_immediate_action', False) else "否"

        recommendations = safety.get('recommendations', [])
        if recommendations and len(recommendations) > 0:
            ws_summary['A13'] = "评估建议"
            ws_summary['A13'].font = Font(bold=True)
            for i, rec in enumerate(recommendations, start=14):
                ws_summary[f'A{i}'] = f"{i-13}. {str(rec)}"

        for col in ['A', 'B']:
            ws_summary.column_dimensions[col].width = 40

        ws_thickness = wb.create_sheet("厚度分析")
        thickness_stats = analysis_results.get('thickness_stats', {}) if analysis_results else {}
        if thickness_stats:
            headers = ["指标", "数值", "单位"]
            data = [
                ["平均厚度", safe_float(thickness_stats.get('mean_thickness', 0)), "m"],
                ["最小厚度", safe_float(thickness_stats.get('min_thickness', 0)), "m"],
                ["最大厚度", safe_float(thickness_stats.get('max_thickness', 0)), "m"],
                ["中位厚度", safe_float(thickness_stats.get('median_thickness', 0)), "m"],
                ["标准差", safe_float(thickness_stats.get('std_thickness', 0)), "m"],
                ["有效点数", safe_int(thickness_stats.get('valid_points', 0)), ""],
                ["25分位数", safe_float(thickness_stats.get('percentile_25', 0)), "m"],
                ["75分位数", safe_float(thickness_stats.get('percentile_75', 0)), "m"]
            ]
            
            for col, header in enumerate(headers, 1):
                cell = ws_thickness.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align

            for row, row_data in enumerate(data, 2):
                for col, value in enumerate(row_data, 1):
                    cell = ws_thickness.cell(row=row, column=col, value=value)
                    cell.font = normal_font
                    cell.alignment = center_align

        ws_thickness.column_dimensions['A'].width = 20
        ws_thickness.column_dimensions['B'].width = 20
        ws_thickness.column_dimensions['C'].width = 10

        ws_wear = wb.create_sheet("磨损分析")
        wear_metrics = analysis_results.get('wear_metrics', {}) if analysis_results else {}
        if wear_metrics:
            headers = ["指标", "数值", "单位"]
            data = [
                ["磨损区域数量", safe_int(wear_metrics.get('num_wear_regions', 0)), ""],
                ["最大磨损深度", abs(safe_float(wear_metrics.get('max_wear_depth', 0))), "m"],
                ["平均磨损深度", abs(safe_float(wear_metrics.get('mean_wear_depth', 0))), "m"],
                ["磨损面积比例", safe_float(wear_metrics.get('wear_ratio', 0))*100, "%"],
                ["总磨损体积", safe_float(wear_metrics.get('total_wear_volume', 0)), "m³"]
            ]
            
            for col, header in enumerate(headers, 1):
                cell = ws_wear.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align

            for row, row_data in enumerate(data, 2):
                for col, value in enumerate(row_data, 1):
                    cell = ws_wear.cell(row=row, column=col, value=value)
                    cell.font = normal_font
                    cell.alignment = center_align

            regions = wear_metrics.get('region_metrics', [])
            if regions and len(regions) > 0:
                ws_regions = wb.create_sheet("磨损区域详情")
                region_headers = ["区域ID", "点数", "最大深度(m)", "平均深度(m)", "估计体积(m³)"]
                
                for col, header in enumerate(region_headers, 1):
                    cell = ws_regions.cell(row=1, column=col, value=header)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center_align

                for row, region in enumerate(regions, 2):
                    ws_regions.cell(row=row, column=1, value=str(region.get('region_id', 'N/A')))
                    ws_regions.cell(row=row, column=2, value=safe_int(region.get('num_points', 0)))
                    ws_regions.cell(row=row, column=3, value=abs(safe_float(region.get('max_depth', 0))))
                    ws_regions.cell(row=row, column=4, value=abs(safe_float(region.get('mean_depth', 0))))
                    ws_regions.cell(row=row, column=5, value=safe_float(region.get('volume', 0)))

                for col in range(1, 6):
                    ws_regions.column_dimensions[chr(64+col)].width = 18

        ws_wear.column_dimensions['A'].width = 20
        ws_wear.column_dimensions['B'].width = 20
        ws_wear.column_dimensions['C'].width = 10

        try:
            wb.save(excel_path)
        except Exception as e:
            print(f"Excel保存警告: {e}")
            pass

        return excel_path

    def _generate_json_report(self, analysis_results: Dict[str, Any],
                              component_id: str, inspector: str,
                              base_filename: str) -> str:
        json_path = os.path.join(self.output_dir, f"{base_filename}.json")

        def convert_enum(obj):
            if hasattr(obj, 'value'):
                return obj.value
            elif isinstance(obj, np.ndarray):
                return obj.tolist()
            elif isinstance(obj, np.generic):
                return obj.item()
            return str(obj)

        report_data = {
            'report_info': {
                'component_id': component_id,
                'inspector': inspector,
                'report_date': datetime.now().isoformat(),
                'version': '1.0'
            },
            'analysis_results': analysis_results
        }

        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(report_data, f, ensure_ascii=False, indent=2, default=convert_enum)

        return json_path

    def generate_summary_dataframe(self, analysis_results: Dict[str, Any]) -> pd.DataFrame:
        data = {
            'Metric': [],
            'Value': [],
            'Unit': [],
            'Category': []
        }

        thickness_stats = analysis_results.get('thickness_stats', {})
        for key, value in thickness_stats.items():
            if key in ['mean_thickness', 'min_thickness', 'max_thickness', 'median_thickness']:
                data['Metric'].append(key)
                data['Value'].append(value)
                data['Unit'].append('m')
                data['Category'].append('Thickness')

        wear_metrics = analysis_results.get('wear_metrics', {})
        for key in ['num_wear_regions', 'max_wear_depth', 'mean_wear_depth', 'wear_ratio']:
            data['Metric'].append(key)
            data['Value'].append(wear_metrics.get(key, 0))
            data['Unit'].append('' if key == 'num_wear_regions' else ('%' if key == 'wear_ratio' else 'm'))
            data['Category'].append('Wear')

        safety = analysis_results.get('safety_summary', {})
        data['Metric'].append('safety_score')
        data['Value'].append(safety.get('safety_score', 0))
        data['Unit'].append('/100')
        data['Category'].append('Safety')

        return pd.DataFrame(data)


__all__ = ['ReportGenerator']
